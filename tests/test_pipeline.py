# -*- coding: utf-8 -*-
"""流水线集成测试：验证各阶段的编排、产物落盘与缓存复用。

设计思路（补上《测试与代码质量报告》中承认的覆盖缺口）：
- 全部离线运行，不打真实网络请求；
- 通过 monkeypatch 把 config 的路径重定向到临时目录，避免污染真实产物；
- 手工写入 raw 缓存文件，验证"跳过 crawl 阶段也能从缓存继续"这一关键能力。
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from llm_danmaku import config, pipeline
from llm_danmaku.cleaner import CleanReport
from llm_danmaku.danmaku_crawler import Danmaku, DanmakuCrawler
from llm_danmaku.video_search import VideoInfo

# ------------------------------------------------------------------ 夹具
SAMPLE_VIDEOS = [
    VideoInfo(bvid="BV1TEST00001", title="大模型科普", author="UP甲", play=1000, rank=1),
    VideoInfo(bvid="BV1TEST00002", title="多模态入门", author="UP乙", play=500, rank=2),
]

SAMPLE_TEXT = {
    "BV1TEST00001": [
        "用GPT写代码效率真高", "大模型会不会取代程序员", "这个AI收费太贵了",
        "用GPT写代码效率真高", "666", "哈哈哈哈哈",
    ],
    "BV1TEST00002": [
        "多模态能力惊艳", "担心隐私泄露问题", "画图效果太好了", "前排打卡",
    ],
}


@pytest.fixture()
def sandbox(tmp_path, monkeypatch):
    """把 config 的所有输出路径重定向到临时目录。"""
    paths = {
        "DATA_DIR": tmp_path / "data",
        "RAW_DIR": tmp_path / "data" / "raw",
        "OUTPUT_DIR": tmp_path / "output",
        "LOG_DIR": tmp_path / "logs",
        "CHART_DIR": tmp_path / "output" / "charts",
        "PROFILE_DIR": tmp_path / "output" / "profile",
        "VIDEO_LIST_FILE": tmp_path / "data" / "videos.json",
        "CLEANED_FILE": tmp_path / "data" / "danmaku_cleaned.jsonl",
        "STATS_FILE": tmp_path / "output" / "danmaku_stats.json",
        "EXCEL_FILE": tmp_path / "output" / "danmaku_analysis.xlsx",
        "WORDCLOUD_FILE": tmp_path / "output" / "wordcloud.png",
        "DASHBOARD_FILE": tmp_path / "output" / "dashboard.html",
    }
    for name, value in paths.items():
        monkeypatch.setattr(config, name, value, raising=True)
    for key in ("DATA_DIR", "RAW_DIR", "OUTPUT_DIR", "LOG_DIR", "CHART_DIR", "PROFILE_DIR"):
        paths[key].mkdir(parents=True, exist_ok=True)

    # 写入视频清单
    with open(paths["VIDEO_LIST_FILE"], "w", encoding="utf-8") as fh:
        json.dump([v.as_dict() for v in SAMPLE_VIDEOS], fh, ensure_ascii=False)

    # 写入 raw 弹幕缓存（模拟已抓取完成的状态）
    crawler = DanmakuCrawler.__new__(DanmakuCrawler)     # 不跑 __init__，只用其写盘方法
    for video in SAMPLE_VIDEOS:
        result = type("R", (), {})()
        from llm_danmaku.danmaku_crawler import CrawlResult
        result = CrawlResult(
            bvid=video.bvid, title=video.title, cid=1000 + video.rank,
            danmaku=[Danmaku(t, progress=i * 60_000) for i, t in
                     enumerate(SAMPLE_TEXT[video.bvid])],
            source="xml", ok=True,
        )
        crawler._save_cache(paths["RAW_DIR"] / f"{video.bvid}.jsonl", result)

    return paths


# ---------------------------------------------------------------- TC50
def test_load_videos_读取清单(sandbox):
    """TC50：视频清单应能完整读回，字段不丢失。"""
    videos = pipeline.load_videos()
    assert len(videos) == 2
    assert videos[0].bvid == "BV1TEST00001"
    assert videos[0].author == "UP甲"


# ---------------------------------------------------------------- TC51
def test_load_videos_文件缺失抛异常(sandbox, monkeypatch):
    """TC51：清单缺失时应抛出带指引的异常，而不是返回空列表。"""
    monkeypatch.setattr(config, "VIDEO_LIST_FILE", sandbox["DATA_DIR"] / "not_exist.json")
    with pytest.raises(FileNotFoundError) as exc:
        pipeline.load_videos()
    assert "search" in str(exc.value)


# ---------------------------------------------------------------- TC52
def test_load_cached_results_复用缓存(sandbox):
    """TC52：跳过 crawl 阶段时，应从 raw 缓存读回弹幕（断点续爬的关键）。"""
    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    assert len(results) == 2
    assert sum(r.count for r in results) == 10          # 6 + 4
    assert results[0].danmaku[0].text == "用GPT写代码效率真高"
    assert results[0].danmaku[0].progress == 0
    assert results[0].danmaku[1].progress == 60_000


# ---------------------------------------------------------------- TC53
def test_load_cached_results_缺缓存时跳过(sandbox):
    """TC53：部分视频没有缓存时应跳过并记录，不抛异常。"""
    extra = VideoInfo(bvid="BV1MISSING001", title="没有缓存", rank=3)
    results = pipeline.load_cached_results([*SAMPLE_VIDEOS, extra])
    assert len(results) == 2                            # 缺缓存的那个被跳过


# ---------------------------------------------------------------- TC54
def test_stage_clean_产物落盘(sandbox):
    """TC54：清洗阶段应生成全量结果与按视频切分的文件。"""
    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    cleaned, report = pipeline.stage_clean(results)

    assert config.CLEANED_FILE.exists()
    lines = config.CLEANED_FILE.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == len(cleaned)
    assert json.loads(lines[0])["text"] == cleaned[0]

    # 噪声与重复应被剔除：原始 10 条 -> 保留 6 条
    assert report.total == 10
    assert report.dropped_noise >= 2                    # 666 / 哈哈哈哈哈 / 前排打卡
    assert report.dropped_duplicate == 1                # 重复的那条

    per_video = config.DATA_DIR / "cleaned_by_video"
    assert (per_video / "BV1TEST00001.txt").exists()
    first_video = (per_video / "BV1TEST00001.txt").read_text(encoding="utf-8")
    assert "666" not in first_video


# ---------------------------------------------------------------- TC55
def test_stage_analyze_写JSON(sandbox):
    """TC55：分析阶段应把统计结果写入 JSON，且字段可读回。"""
    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    cleaned, report = pipeline.stage_clean(results)
    stats = pipeline.stage_analyze(results, cleaned, report)

    assert config.STATS_FILE.exists()
    raw = json.loads(config.STATS_FILE.read_text(encoding="utf-8"))
    assert raw["video_count"] == 2
    assert raw["total_raw"] == 10
    assert raw["total_cleaned"] == len(cleaned)
    assert raw["top_comments"]
    assert stats.top_words


# ---------------------------------------------------------------- TC56
def test_load_stats_还原tuple(sandbox):
    """TC56：load_stats 应把 JSON 里的 list 还原成 tuple（供图表模块使用）。"""
    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    cleaned, report = pipeline.stage_clean(results)
    pipeline.stage_analyze(results, cleaned, report)

    stats = pipeline.load_stats()
    assert stats.video_count == 2
    assert isinstance(stats.top_comments[0], tuple)
    assert isinstance(stats.video_stats[0].top_comments[0], tuple)
    assert isinstance(stats.progress_buckets[0], tuple)
    assert stats.domain_counts.get("编程开发", 0) >= 1
    assert stats.cost_count >= 1


# ---------------------------------------------------------------- TC57
def test_load_stats_文件缺失(sandbox, monkeypatch):
    """TC57：统计文件缺失时应抛出带指引的异常。"""
    monkeypatch.setattr(config, "STATS_FILE", sandbox["OUTPUT_DIR"] / "none.json")
    with pytest.raises(FileNotFoundError) as exc:
        pipeline.load_stats()
    assert "analyze" in str(exc.value)


# ---------------------------------------------------------------- TC58
def test_stage_export_生成Excel(sandbox):
    """TC58：导出阶段应生成 xlsx 并包含约定的 Sheet。"""
    openpyxl = pytest.importorskip("openpyxl")     # noqa: F841
    from openpyxl import load_workbook

    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    cleaned, report = pipeline.stage_clean(results)
    stats = pipeline.stage_analyze(results, cleaned, report)
    path = pipeline.stage_export(stats)

    assert Path(path).exists()
    book = load_workbook(path)
    for sheet in ["数据总览", "Top弹幕", "词频统计", "视频清单", "领域分布", "情感与立场"]:
        assert sheet in book.sheetnames


# ---------------------------------------------------------------- TC59
def test_stage_visualize_出图(sandbox):
    """TC59：可视化阶段应产出词云与图表文件（依赖 matplotlib/wordcloud）。"""
    pytest.importorskip("matplotlib")
    pytest.importorskip("wordcloud")

    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    cleaned, report = pipeline.stage_clean(results)
    stats = pipeline.stage_analyze(results, cleaned, report)
    charts = pipeline.stage_visualize(stats)

    assert charts.get("wordcloud") and Path(charts["wordcloud"]).exists()
    assert Path(charts["wordcloud"]).stat().st_size > 1000      # 不是空图
    assert charts.get("top_comments") and Path(charts["top_comments"]).exists()


# ---------------------------------------------------------------- TC60
def test_stage_dashboard_出HTML(sandbox):
    """TC60：大屏阶段应产出单文件 HTML，且包含指标卡片。"""
    pytest.importorskip("pyecharts")

    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    cleaned, report = pipeline.stage_clean(results)
    stats = pipeline.stage_analyze(results, cleaned, report)
    path = pipeline.stage_dashboard(stats)

    assert path is not None and Path(path).exists()
    html = Path(path).read_text(encoding="utf-8")
    assert "kpi-card" in html                     # 注入的关键指标卡片
    assert "抓取视频" in html


# ---------------------------------------------------------------- TC61
def test_run_pipeline_未知阶段报错(sandbox):
    """TC61：未知阶段名应抛出 ValueError（CLI 会转成退出码 2）。"""
    with pytest.raises(ValueError) as exc:
        pipeline.run_pipeline(["不存在的阶段"])
    assert "未知阶段" in str(exc.value)


# ---------------------------------------------------------------- TC62
def test_main_未知阶段返回2(sandbox):
    """TC62：CLI 对非法参数应返回退出码 2，而不是崩溃。"""
    assert pipeline.main(["--stage", "不存在的阶段"]) == 2
    assert pipeline.main(["--stages", "clean,bad,analyze"]) == 2


# ---------------------------------------------------------------- TC63
def test_main_离线跑通数据链路(sandbox):
    """TC63：CLI 在离线（无 crawl）情况下应能跑通 clean/analyze/export 全链路。"""
    pytest.importorskip("matplotlib")
    code = pipeline.main(["--stages", "clean,analyze,export"])
    assert code == 0
    assert config.EXCEL_FILE.exists()
    assert config.STATS_FILE.exists()


# ---------------------------------------------------------------- TC64
def test_setup_logging_可重复调用(sandbox):
    """TC64：日志初始化应可重复调用且不抛异常（force=True 分支）。"""
    pipeline.setup_logging("WARNING")
    pipeline.setup_logging("INFO")
    assert (config.LOG_DIR / "run.log").exists()


# ---------------------------------------------------------------- TC65
def test_load_cleaned_读回(sandbox):
    """TC65：清洗结果应能从落盘文件读回，报告计数与实际一致。"""
    results = pipeline.load_cached_results(SAMPLE_VIDEOS)
    cleaned, _ = pipeline.stage_clean(results)
    loaded, report = pipeline.load_cleaned()
    assert loaded == cleaned
    assert report.total == len(cleaned)
    assert isinstance(report, CleanReport)
